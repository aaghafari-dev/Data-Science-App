"""LangGraph Presentation Agent with evidence-first slide planning and professional narrative."""
from __future__ import annotations
import json, os, tempfile
from datetime import datetime
from pathlib import Path
from typing import Any, TypedDict
from langchain_core.messages import AIMessage, HumanMessage
from langgraph.graph import END, StateGraph
from services.agent_memory import AgentMemory
from services.llm_config import provider_from_state, get_llm
from services.serialization import json_safe
from services.narrative_engine import EvidenceNarrativeEngine
from services.output_quality import OutputQualityGate

class PresentationState(TypedDict, total=False):
    report_evidence: dict[str,Any]; report_text: str; plot_results: dict[str,Any]; memory:list[dict[str,Any]]
    output_path:str; title:str; presenter:str; status:str; summary:str; narrative:dict[str,str]; slides:list[dict[str,Any]]
    error:str; llm_provider:str; model_name:str; api_key:str; local_path:str; api_base:str

def _plan(state):
    """Plan a presentation around one analytical question per slide."""
    evidence=json_safe(state.get("report_evidence") or {}); plots=json_safe(state.get("plot_results") or evidence.get("Plot Agent") or {})
    plot_items=(plots or {}).get("plots",[]) if isinstance(plots,dict) else []
    title=state.get("title") or "Data Science Analysis Presentation"; cfg=provider_from_state(state); llm=get_llm(cfg)
    narrative=EvidenceNarrativeEngine.deterministic(evidence)
    if llm:
        prompt=EvidenceNarrativeEngine.build_prompt(evidence, audience="senior data-science colleagues", format_name="executive analytical presentation")
        try:
            r=llm.invoke([HumanMessage(content=prompt)]); raw=r.content if hasattr(r,"content") else str(r); m=__import__("re").search(r"\{.*\}",raw,__import__("re").S); obj=json.loads(m.group(0)) if m else {}
            if isinstance(obj,dict):
                for k,v in obj.items():
                    if isinstance(v,str) and v.strip(): narrative[k]=v.strip()
        except Exception: pass
    slides=[
        {"title":title,"purpose":"Orient the audience to the question and scope.","body":EvidenceNarrativeEngine.slide_text(narrative["executive_summary"]),"notes":narrative["executive_summary"]},
        {"title":"Analytical Question and Data","purpose":"Explain what was analysed and what evidence was available.","body":EvidenceNarrativeEngine.slide_text(narrative["objective"]+" "+narrative["data_governance"]),"notes":narrative["data_governance"]},
        {"title":"Why This Analytical Strategy?","purpose":"Show how task identification and method selection were governed.","body":EvidenceNarrativeEngine.slide_text(narrative["methodology"]),"notes":narrative["methodology"]},
        {"title":"Validation Strategy","purpose":"Explain how generalisation or stability was assessed.","body":EvidenceNarrativeEngine.slide_text(narrative["validation"]),"notes":narrative["validation"]},
        {"title":"Key Results","purpose":"Present the principal quantitative evidence.","body":EvidenceNarrativeEngine.slide_text(narrative["results"]),"notes":narrative["results"]},
        {"title":"Robustness and Interpretation","purpose":"Separate observed performance from robustness and interpretation.","body":EvidenceNarrativeEngine.slide_text(narrative["robustness"]+" "+narrative["discussion"]),"notes":narrative["discussion"]},
        {"title":"Limitations and Next Steps","purpose":"Make uncertainty and remaining work explicit.","body":EvidenceNarrativeEngine.slide_text(narrative["limitations"]+" "+narrative["recommendations"]),"notes":narrative["recommendations"]},
        {"title":"Conclusion","purpose":"Close with evidence-bound conclusions.","body":EvidenceNarrativeEngine.slide_text(narrative["conclusion"]),"notes":narrative["conclusion"]},
    ]
    for i,item in enumerate(plot_items[:6]):
        slides.insert(5+i,{"title":item.get("title",f"Evidence Visualization {i+1}"),"purpose":"Show one evidence-linked visual finding.","body":EvidenceNarrativeEngine.slide_text(" ".join(map(str,item.get("insights",[]))) or str(item.get("rationale") or "Evidence-linked visualization.")),"notes":item.get("rationale") or "Interpret the visualization together with its underlying evidence." ,"figure_json":item.get("figure_json")})
    quality=OutputQualityGate.presentation(slides,evidence)
    mem=AgentMemory("Presentation",30); mem.remember("presentation_plan",{"title":title,"slide_count":len(slides),"plot_count":len(plot_items),"provider":cfg.provider,"quality":quality})
    return {"status":"planned","slides":slides,"output_quality":quality,"narrative":narrative,"memory":mem.to_dict(),"messages":[AIMessage(content=f"Planned a coherent {len(slides)}-slide evidence-first presentation.")]}

def _plot_images(slides):
    """Render evidence-linked Plotly figures for presentation slides."""
    images=[]
    try: import plotly.io as pio
    except Exception: return images
    for i,slide in enumerate(slides):
        fj=slide.get("figure_json")
        if not fj: continue
        path=Path(tempfile.gettempdir())/f"dssp_presentation_plot_{os.getpid()}_{i}.png"
        try:
            pio.from_json(fj).write_image(str(path),format="png",width=1400,height=760,scale=1)
            if path.exists(): images.append((slide,path))
        except Exception: pass
    return images

def _build(state):
    """Build a polished PowerPoint with concise slide text and speaker notes."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except Exception as exc: return {"status":"error","error":f"PowerPoint generation requires python-pptx: {exc}"}
    slides=json_safe(state.get("slides") or []); title=state.get("title") or "Data Science Analysis Presentation"; presenter=state.get("presenter") or "Data Science Studio Pro"
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5); blank=prs.slide_layouts[6]
    rendered=dict(_plot_images(slides))
    for idx,plan in enumerate(slides,1):
        slide=prs.slides.add_slide(blank)
        tb=slide.shapes.add_textbox(Inches(.65),Inches(.35),Inches(12),Inches(.75)); p=tb.text_frame.paragraphs[0]; p.text=plan.get("title",f"Slide {idx}"); p.font.size=Pt(28); p.font.bold=True
        sub=slide.shapes.add_textbox(Inches(.7),Inches(1.05),Inches(11.8),Inches(.4)); p=sub.text_frame.paragraphs[0]; p.text=plan.get("purpose",""); p.font.size=Pt(12)
        if idx==1:
            body=f"Prepared by: {presenter}\n{datetime.now():%Y-%m-%d}\n\n{plan.get('body','')}"
        else: body=plan.get("body","")
        if plan.get("figure_json") and plan.get("title") in rendered:
            slide.shapes.add_picture(str(rendered[plan["title"]]),Inches(.7),Inches(1.55),width=Inches(8.4),height=Inches(4.9))
            box=slide.shapes.add_textbox(Inches(9.35),Inches(1.7),Inches(3.2),Inches(4.6)); tf=box.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.text=body; p.font.size=Pt(16)
        else:
            box=slide.shapes.add_textbox(Inches(.9),Inches(1.65),Inches(11.4),Inches(4.8)); tf=box.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.text=body; p.font.size=Pt(20)
        footer=slide.shapes.add_textbox(Inches(.7),Inches(7.05),Inches(11.8),Inches(.25)); p=footer.text_frame.paragraphs[0]; p.text=f"Data Science Studio Pro • Evidence-bound analysis • {idx}/{len(slides)}"; p.font.size=Pt(9)
        try:
            notes=slide.notes_slide.notes_text_frame; notes.text=plan.get("notes",body)
        except Exception: pass
    output=state.get("output_path") or "DataScienceStudioPro_Analysis_Presentation.pptx"; Path(output).parent.mkdir(parents=True,exist_ok=True); prs.save(output)
    for _,path in rendered.items(): path.unlink(missing_ok=True)
    return {"status":"complete","output_path":output,"summary":f"Created a {len(slides)}-slide professional presentation with one analytical message per slide and speaker notes.","messages":[AIMessage(content=f"Generated {output}")]}

workflow=StateGraph(PresentationState); workflow.add_node("plan",_plan); workflow.add_node("build",_build); workflow.set_entry_point("plan"); workflow.add_edge("plan","build"); workflow.add_edge("build",END); agent_presentation_app=workflow.compile()
